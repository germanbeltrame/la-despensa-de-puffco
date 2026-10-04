"""Corrige costos de Stock y price: si M > H se usa H; si H coincide con M/N se extrae el FOB."""

import os
import re
import sys
import unicodedata
from pathlib import Path

from dotenv import load_dotenv
from openpyxl import load_workbook
from supabase import create_client

EXCEL = Path(__file__).resolve().parent / "PUFFCO APP.xlsx"
HOJA = "Stock y price"
TARIFA_KG = 38.0
FACTOR = 1.065
ENCABEZADOS = {"producto", "producto (a)", "nombre"}


def configurar_consola() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")


def texto(valor) -> str:
    if valor is None:
        return ""
    return " ".join(str(valor).replace("\xa0", " ").split())


def numero(valor):
    if isinstance(valor, bool) or valor is None:
        return None
    if isinstance(valor, (int, float)):
        return float(valor)
    return None


def clave_nombre(nombre: str) -> str:
    base = unicodedata.normalize("NFD", nombre)
    sin_acentos = "".join(letra for letra in base if unicodedata.category(letra) != "Mn")
    return " ".join(sin_acentos.casefold().split())


def sumando(formula) -> float | None:
    coincidencia = re.fullmatch(r"=L\d+\+([0-9]+(?:\.[0-9]+)?)", str(formula or "").replace(" ", ""))
    if not coincidencia:
        return None
    return float(coincidencia.group(1))


def leer_costos() -> list[dict]:
    formulas = list(load_workbook(EXCEL, data_only=False, read_only=True)[HOJA].iter_rows(min_row=2, max_col=15))
    valores = list(load_workbook(EXCEL, data_only=True, read_only=True)[HOJA].iter_rows(min_row=2, max_col=15))
    productos = []
    vistos = set()
    for ff, vv in zip(formulas, valores):
        nombre = texto(ff[0].value)
        if not nombre or clave_nombre(nombre) in ENCABEZADOS:
            continue
        if not any(letra.isalpha() for letra in nombre):
            continue
        clave = clave_nombre(nombre)
        if clave in vistos:
            continue
        vistos.add(clave)
        peso = numero(vv[14].value)
        costo_h = numero(vv[7].value)
        fob_m = sumando(ff[12].value)
        peso_kg = 0.0 if peso is None else round(peso, 3)
        if costo_h is None or fob_m is None:
            productos.append(_directo(nombre, vv, peso_kg, None, False))
            continue
        aereo = peso_kg * TARIFA_KG
        columna_m = aereo + fob_m
        columna_n = columna_m * FACTOR
        if columna_m > costo_h + 0.01:
            productos.append(_directo(nombre, vv, peso_kg, costo_h, False))
            continue
        tolerancia = max(0.25, abs(costo_h) * 0.005)
        coherente = abs(costo_h - columna_n) <= tolerancia
        productos.append(_directo(nombre, vv, peso_kg, costo_h, coherente, fob_m if coherente else None))
    return productos


def _directo(nombre, fila, peso_kg, costo_h, fob_verificado, fob=None) -> dict:
    return {
        "nombre": nombre,
        "precio_puntero_usd": round(numero(fila[1].value) or 0, 2),
        "stock_actual": max(int(numero(fila[3].value) or 0), 0),
        "marca": texto(fila[5].value) or "PUFFCO",
        "categoria": texto(fila[6].value) or "VAPORIZADOR",
        "peso_kg": peso_kg,
        "precio_kg": TARIFA_KG,
        "costo_fob": round(fob, 2) if fob_verificado and fob is not None else 0.0,
        "costo_total_usd": None if costo_h is None else round(costo_h, 2),
        "calcular_costo": False,
        "fob_verificado": fob_verificado and (fob or 0) > 0,
        "regla": "fob" if fob_verificado else "h",
    }


def flete_de_38(sb) -> int:
    filas = sb.table("tipos_flete").select("id,costo_usd_kg").execute().data or []
    for fila in filas:
        if abs(float(fila.get("costo_usd_kg") or 0) - TARIFA_KG) < 0.001:
            return int(fila["id"])
    creado = (
        sb.table("tipos_flete")
        .insert({"nombre": "USD 38.00/kg", "costo_usd_kg": TARIFA_KG})
        .execute()
        .data
        or []
    )
    if creado:
        return int(creado[0]["id"])
    raise RuntimeError("No se pudo guardar la tarifa de US$ 38 por kg.")


def existentes(sb) -> dict[str, dict]:
    encontrados = {}
    inicio = 0
    while True:
        lote = (
            sb.table("productos")
            .select("id,nombre")
            .range(inicio, inicio + 999)
            .execute()
            .data
            or []
        )
        for fila in lote:
            encontrados[clave_nombre(fila["nombre"])] = fila
        if len(lote) < 1000:
            return encontrados
        inicio += 1000


def aplicar(sb, productos: list[dict], flete_id: int) -> tuple[int, int]:
    cargados = existentes(sb)
    actualizados = 0
    nuevos = 0
    for producto in productos:
        payload = {
            "peso_kg": producto["peso_kg"],
            "precio_kg": producto["precio_kg"],
            "flete_id": flete_id,
            "costo_fob": producto["costo_fob"],
            "costo_total_usd": producto["costo_total_usd"],
            "calcular_costo": False,
        }
        actual = cargados.get(clave_nombre(producto["nombre"]))
        if actual is None:
            sb.table("productos").insert(
                {
                    **payload,
                    "nombre": producto["nombre"],
                    "marca": producto["marca"],
                    "categoria": producto["categoria"],
                    "precio_puntero_usd": producto["precio_puntero_usd"],
                    "stock_actual": producto["stock_actual"],
                }
            ).execute()
            nuevos += 1
            continue
        sb.table("productos").update(payload).eq("id", int(actual["id"])).execute()
        actualizados += 1
    return actualizados, nuevos


def main() -> None:
    configurar_consola()
    if not EXCEL.exists():
        raise FileNotFoundError(f"No se encontró el Excel: {EXCEL}")
    load_dotenv(Path(__file__).resolve().parent / ".env", override=True)
    url = os.getenv("SUPABASE_URL", "").strip().strip('"')
    clave = os.getenv("SUPABASE_KEY", "").strip().strip('"')
    if not url or not clave:
        raise RuntimeError("Faltan SUPABASE_URL o SUPABASE_KEY en el archivo .env.")

    productos = leer_costos()
    con_fob = [item for item in productos if item["fob_verificado"]]
    directos = [item for item in productos if not item["fob_verificado"]]
    if len(con_fob) < 100:
        raise RuntimeError(f"Solo {len(con_fob)} FOB coherentes. No se escribió nada.")

    sb = create_client(url, clave)
    flete_id = flete_de_38(sb)
    actualizados, nuevos = aplicar(sb, productos, flete_id)
    incompletos = sum(1 for item in productos if item["peso_kg"] == 0 or not item["fob_verificado"])

    print(f"Tarifa US$ 38/kg, flete id {flete_id}")
    print(f"Productos de la solapa: {len(productos)}")
    print(f"FOB extraído de M/N: {len(con_fob)}")
    print(f"Costo directo de H (M mayor que H, o fórmula no coherente): {len(directos)}")
    print(f"Incompletos (sin peso o sin FOB verificado): {incompletos}")
    print(f"Actualizados: {actualizados}  |  Nuevos: {nuevos}")
    print()
    print("Ejemplos con FOB verificado:")
    for item in con_fob[:3]:
        print(f"  {item['nombre']}  |  FOB {item['costo_fob']:.2f}  |  costo H {item['costo_total_usd']:.2f}")
    print("Ejemplos con costo directo de H:")
    for item in directos[:5]:
        costo = "s/d" if item["costo_total_usd"] is None else f"{item['costo_total_usd']:.2f}"
        print(f"  {item['nombre']}  |  costo H {costo}  |  peso {item['peso_kg']}")


if __name__ == "__main__":
    main()
